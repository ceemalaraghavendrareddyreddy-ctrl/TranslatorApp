"""
Translator App
--------------
Translates French, Bengali, Hindi, Malagasy, Nepali <-> English.
Supports Excel (.xlsx), Word (.docx) and PDF (.pdf) input, preserving
original layout as much as possible. Free translation via deep-translator
(no API key / no billing). Generates a translated output file of the same
type, and lets you send it straight to the printer.
"""

import os
import sys
import copy
import time
import traceback
from concurrent.futures import ThreadPoolExecutor

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QLabel,
    QPushButton, QFileDialog, QComboBox, QProgressBar, QMessageBox,
    QTextEdit, QGroupBox, QToolButton, QTabWidget, QCheckBox
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtPrintSupport import QPrinter, QPrintDialog

import requests
from deep_translator import GoogleTranslator, MyMemoryTranslator

import openpyxl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pdfplumber
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

UNICODE_FONT_NAME = "Helvetica"
_UNICODE_FONT_CANDIDATES = [
    "C:/Windows/Fonts/ARIALUNI.TTF",
    "C:/Windows/Fonts/Nirmala.ttf",
    "C:/Windows/Fonts/arial.ttf",
]
for _font_path in _UNICODE_FONT_CANDIDATES:
    if os.path.exists(_font_path):
        try:
            pdfmetrics.registerFont(TTFont("UnicodeFont", _font_path))
            UNICODE_FONT_NAME = "UnicodeFont"
            break
        except Exception:
            continue

DOCX_UNICODE_FONT = "Nirmala UI"  # covers Devanagari (Hindi/Nepali), Bengali, and Latin scripts


def set_run_font(run, font_name=DOCX_UNICODE_FONT):
    """Force a run to use a font that actually has glyphs for Hindi/Bengali/
    Nepali script (Word's default document font usually doesn't, which shows
    as empty boxes)."""
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)


LANGUAGES = {
    "English": "en",
    "French": "fr",
    "Bengali": "bn",
    "Hindi": "hi",
    "Malagasy": "mg",
    "Nepali": "ne",
}
CODE_TO_NAME = {v: k for k, v in LANGUAGES.items()}

# MyMemoryTranslator (the fallback free engine) needs locale-style codes
# (e.g. "mg-MG"), not the plain ISO codes GoogleTranslator accepts.
MYMEMORY_CODES = {
    "en": "en-GB",
    "fr": "fr-FR",
    "bn": "bn-IN",
    "hi": "hi-IN",
    "mg": "mg-MG",
    "ne": "ne-NP",
}

MAX_CHARS = 4500  # deep_translator/Google limit safety margin

# Google's free translate endpoint has no official quota; hammering it with
# back-to-back requests trips its rate limiting (worse the more it's used,
# and worse on some days than others). A short delay between successful
# calls plus a capped exponential backoff on failures rides out most
# throttling windows without letting one stuck segment stall the whole
# document -- the backoff is capped so a bad run costs seconds, not
# tens of seconds, per failing segment.
REQUEST_DELAY = 0.15        # seconds to wait after a successful translate call
MAX_RETRIES = 3             # attempts per piece of text before giving up
RETRY_BACKOFF_BASE = 1.5    # seconds; backoff is RETRY_BACKOFF_BASE * 2**attempt
RETRY_BACKOFF_CAP = 6.0     # seconds; upper bound on any single backoff wait

# If Google fails this many segments in a row, it's not a fluke -- it's
# actively rate-limiting this connection for the whole run. Retrying it on
# every subsequent segment would just waste ~MAX_RETRIES worth of backoff
# each time, so once this threshold is hit, skip straight to the fallback
# engine for the rest of the document instead of re-proving Google is down.
GOOGLE_SKIP_AFTER_CONSECUTIVE_FAILURES = 3

# Optional paid fallback: if the free scrape exhausts every retry above,
# and this env var is set to a real Google Cloud Translation API key, the
# app spends one paid call to rescue just that segment instead of leaving
# it untranslated. Leave the env var unset to keep the app 100% free (it
# will simply behave as before -- failed segments stay in the source
# language). Get a key at https://console.cloud.google.com (enable the
# "Cloud Translation API" on a project with billing enabled).
GOOGLE_API_KEY = os.environ.get("GOOGLE_TRANSLATE_API_KEY", "").strip()
GOOGLE_CLOUD_TRANSLATE_URL = "https://translation.googleapis.com/language/translate/v2"


def paid_google_translate(text, source, target):
    """Single call to the official Google Cloud Translation API (REST v2,
    API-key auth). Returns None on any failure so callers can treat it the
    same as a failed free-scrape attempt."""
    if not GOOGLE_API_KEY:
        return None
    try:
        resp = requests.post(
            GOOGLE_CLOUD_TRANSLATE_URL,
            params={"key": GOOGLE_API_KEY},
            data={"q": text, "source": source, "target": target, "format": "text"},
            timeout=15,
        )
        if resp.status_code != 200:
            return None
        data = resp.json()
        return data["data"]["translations"][0]["translatedText"]
    except Exception:
        return None


def chunk_text(text, max_len=MAX_CHARS):
    """Split text into chunks under max_len, breaking on line boundaries."""
    if not text:
        return [text]
    if len(text) <= max_len:
        return [text]
    lines = text.split("\n")
    chunks = []
    current = ""
    for line in lines:
        candidate = (current + "\n" + line) if current else line
        if len(candidate) > max_len:
            if current:
                chunks.append(current)
            if len(line) > max_len:
                for i in range(0, len(line), max_len):
                    chunks.append(line[i:i + max_len])
                current = ""
            else:
                current = line
        else:
            current = candidate
    if current:
        chunks.append(current)
    return chunks


class Translator:
    def __init__(self, source, target):
        self.source = source
        self.target = target
        self._engine = GoogleTranslator(source=source, target=target)
        # Google's free endpoint is the primary engine, but it rate-limits
        # (HTTP 429) independently of anything this app does. MyMemory is a
        # second free, no-key engine -- when Google is throttling, this lets
        # a segment still get translated instead of silently giving up.
        try:
            mm_source = MYMEMORY_CODES.get(source, source)
            mm_target = MYMEMORY_CODES.get(target, target)
            self._fallback_engine = MyMemoryTranslator(source=mm_source, target=mm_target)
        except Exception:
            self._fallback_engine = None
        # Visibility into how many pieces failed after all retries, so the
        # UI can tell the user "N segments were not translated" instead of
        # silently handing back a partly-English document.
        self.total_count = 0
        self.fail_count = 0
        self.fallback_engine_count = 0
        self.paid_fallback_count = 0
        self.cache_hits = 0
        # Same text (repeated headers, labels, boilerplate) shows up many
        # times in real documents; translating it again each time wastes a
        # network round trip for a result we already have.
        self._cache = {}
        # Once Google proves it's rate-limiting this run, stop wasting
        # MAX_RETRIES worth of backoff on it for every remaining segment.
        self._consecutive_google_failures = 0
        self.google_skipped = False

    def _translate_once(self, text):
        """Translate a single piece of text, retrying with capped backoff
        on failure. If Google's free engine fails every retry, try the
        MyMemory free engine as a fallback. If that also fails and a paid
        API key is configured, spend one paid call to rescue just this
        piece. Returns None if everything failed."""
        if not self.google_skipped:
            for attempt in range(MAX_RETRIES):
                try:
                    result = self._engine.translate(text)
                except Exception:
                    result = None
                if result is not None:
                    time.sleep(REQUEST_DELAY)
                    self._consecutive_google_failures = 0
                    return result
                if attempt < MAX_RETRIES - 1:
                    time.sleep(min(RETRY_BACKOFF_BASE * (2 ** attempt), RETRY_BACKOFF_CAP))

            self._consecutive_google_failures += 1
            if self._consecutive_google_failures >= GOOGLE_SKIP_AFTER_CONSECUTIVE_FAILURES:
                self.google_skipped = True

        if self._fallback_engine is not None:
            for attempt in range(2):
                try:
                    result = self._fallback_engine.translate(text)
                except Exception:
                    result = None
                if result is not None:
                    time.sleep(REQUEST_DELAY)
                    self.fallback_engine_count += 1
                    return result
                time.sleep(min(RETRY_BACKOFF_BASE * (2 ** attempt), RETRY_BACKOFF_CAP))

        if GOOGLE_API_KEY:
            result = paid_google_translate(text, self.source, self.target)
            if result is not None:
                self.paid_fallback_count += 1
                return result
        return None

    def translate(self, text):
        if text is None:
            return text
        stripped = text.strip()
        if not stripped:
            return text

        if text in self._cache:
            self.cache_hits += 1
            return self._cache[text]

        self.total_count += 1
        try:
            if len(text) <= MAX_CHARS:
                result = self._translate_once(text)
                if result is None:
                    self.fail_count += 1
                    result = text
            else:
                parts = chunk_text(text)
                translated_parts = []
                any_failed = False
                for p in parts:
                    part_result = self._translate_once(p)
                    if part_result is None:
                        any_failed = True
                        translated_parts.append(p)
                    else:
                        translated_parts.append(part_result)
                if any_failed:
                    self.fail_count += 1
                result = "\n".join(translated_parts)
        except Exception:
            # Last-resort safety net: keep original text rather than
            # crashing the whole document run.
            self.fail_count += 1
            result = text

        self._cache[text] = result
        return result


def translate_segment_all_languages(original, translators, pool):
    """Translate one piece of text into every language in `translators`
    (code -> Translator) concurrently, using the given thread pool. Each
    language's Translator wraps its own independent engine and keeps its
    own cache/retry state, so there's no shared mutable state between
    threads here -- running them in parallel turns N sequential network
    round-trips (plus N separate REQUEST_DELAY sleeps and N separate retry
    backoffs) into roughly the cost of a single round-trip, instead of
    stacking them one after another."""
    futures = {
        pool.submit(translator.translate, original): code
        for code, translator in translators.items()
    }
    return {code: future.result() for future, code in futures.items()}


def translate_xlsx(in_path, out_path, translator, progress_cb=None):
    wb = openpyxl.load_workbook(in_path)
    cells = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    cells.append(cell)
    total = len(cells) or 1
    for i, cell in enumerate(cells):
        cell.value = translator.translate(cell.value)
        cell.font = openpyxl.styles.Font(name=DOCX_UNICODE_FONT, size=cell.font.size)
        if progress_cb:
            progress_cb(int((i + 1) / total * 100))
    wb.save(out_path)

    preview_lines = []
    for sheet in wb.worksheets:
        preview_lines.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows():
            values = [str(c.value) if c.value is not None else "" for c in row]
            if any(v.strip() for v in values):
                preview_lines.append("\t".join(values))
    stats = (translator.total_count, translator.fail_count, translator.paid_fallback_count, translator.fallback_engine_count, translator.google_skipped)
    return "\n".join(preview_lines), stats


def translate_xlsx_multi(in_path, out_path, translators, progress_cb=None):
    """Each cell becomes: original text, then one line per target language
    labelled with the language name, all stacked inside the same cell."""
    wb = openpyxl.load_workbook(in_path)
    cells = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    cells.append(cell)
    total = len(cells) or 1
    with ThreadPoolExecutor(max_workers=max(1, len(translators))) as pool:
        for i, cell in enumerate(cells):
            original = cell.value
            translated_map = translate_segment_all_languages(original, translators, pool)
            lines = [original]
            for code in translators:
                lines.append(f"[{CODE_TO_NAME[code]}] {translated_map[code]}")
            cell.value = "\n".join(lines)
            cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
            cell.font = openpyxl.styles.Font(name=DOCX_UNICODE_FONT, size=cell.font.size)
            if progress_cb:
                progress_cb(int((i + 1) / total * 100))
    wb.save(out_path)

    preview_lines = []
    for sheet in wb.worksheets:
        preview_lines.append(f"--- Sheet: {sheet.title} ---")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip():
                    preview_lines.append(cell.value)
                    preview_lines.append("")
    stats = (
        sum(t.total_count for t in translators.values()),
        sum(t.fail_count for t in translators.values()),
        sum(t.paid_fallback_count for t in translators.values()),
        sum(t.fallback_engine_count for t in translators.values()),
        any(t.google_skipped for t in translators.values()),
    )
    return "\n".join(preview_lines), stats


def convert_doc_to_docx(doc_path):
    """Convert a legacy .doc file to .docx using Word COM automation
    (python-docx cannot read the old binary .doc format directly)."""
    import win32com.client

    doc_path = os.path.abspath(doc_path)
    docx_path = os.path.splitext(doc_path)[0] + "_converted.docx"

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(doc_path)
        doc.SaveAs2(docx_path, FileFormat=16)  # wdFormatXMLDocument (.docx)
        doc.Close()
    finally:
        word.Quit()
    return docx_path


def translate_docx(in_path, out_path, translator, progress_cb=None):
    doc = Document(in_path)

    def translate_paragraph(paragraph):
        if not paragraph.text.strip():
            return
        translated = translator.translate(paragraph.text)
        if paragraph.runs:
            paragraph.runs[0].text = translated
            set_run_font(paragraph.runs[0])
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            run = paragraph.add_run(translated)
            set_run_font(run)

    targets = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)

    total = len(targets) or 1
    for i, paragraph in enumerate(targets):
        translate_paragraph(paragraph)
        if progress_cb:
            progress_cb(int((i + 1) / total * 100))
    doc.save(out_path)

    preview_lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cell_texts = [c.text for c in row.cells]
            if any(t.strip() for t in cell_texts):
                preview_lines.append(" | ".join(cell_texts))
    stats = (translator.total_count, translator.fail_count, translator.paid_fallback_count, translator.fallback_engine_count, translator.google_skipped)
    return "\n".join(preview_lines), stats


def translate_docx_multi(in_path, out_path, translators, progress_cb=None):
    """Each paragraph becomes: original text, then one new line per target
    language (labelled), stacked one after another within the same paragraph."""
    doc = Document(in_path)

    def expand_paragraph(paragraph, pool):
        original = paragraph.text
        if not original.strip():
            return
        if paragraph.runs:
            paragraph.runs[0].text = original
            for run in paragraph.runs[1:]:
                run.text = ""
        translated_map = translate_segment_all_languages(original, translators, pool)
        for code in translators:
            translated = translated_map[code]
            paragraph.add_run().add_break()
            label_run = paragraph.add_run(f"[{CODE_TO_NAME[code]}] ")
            label_run.bold = True
            set_run_font(label_run)
            text_run = paragraph.add_run(translated)
            set_run_font(text_run)

    targets = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                targets.extend(cell.paragraphs)

    total = len(targets) or 1
    with ThreadPoolExecutor(max_workers=max(1, len(translators))) as pool:
        for i, paragraph in enumerate(targets):
            expand_paragraph(paragraph, pool)
            if progress_cb:
                progress_cb(int((i + 1) / total * 100))
    doc.save(out_path)

    preview_lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cell_texts = [c.text for c in row.cells]
            if any(t.strip() for t in cell_texts):
                preview_lines.append(" | ".join(cell_texts))
    stats = (
        sum(t.total_count for t in translators.values()),
        sum(t.fail_count for t in translators.values()),
        sum(t.paid_fallback_count for t in translators.values()),
        sum(t.fallback_engine_count for t in translators.values()),
        any(t.google_skipped for t in translators.values()),
    )
    return "\n".join(preview_lines), stats


def translate_pdf(in_path, out_path, translator, progress_cb=None):
    """Rebuild a new PDF, placing translated text blocks at approximately
    the same position as the original words/lines (layout-preserving best
    effort, not a true PDF text replacement)."""
    preview_lines = []
    with pdfplumber.open(in_path) as pdf:
        # Progress is reported per translated line, not per page -- a PDF
        # with a handful of pages but hundreds of lines would otherwise
        # leave the progress bar sitting at 0% for the entire first page.
        pages_data = []
        total_lines = 0
        for page in pdf.pages:
            lines = [l for l in (page.extract_text_lines() or []) if l.get("text", "").strip()]
            pages_data.append((page.width, page.height, lines))
            total_lines += len(lines)
        total = total_lines or 1

        c = canvas.Canvas(out_path, pagesize=letter)
        done = 0
        for i, (width, height, lines) in enumerate(pages_data):
            c.setPageSize((width, height))
            preview_lines.append(f"--- Page {i + 1} ---")
            for line in lines:
                text = line["text"]
                translated = translator.translate(text) or text
                preview_lines.append(translated)
                x = line["x0"]
                y = height - line["top"] - (line["bottom"] - line["top"])
                font_size = max(6, min(14, line["bottom"] - line["top"]))
                c.setFont(UNICODE_FONT_NAME, font_size)
                c.drawString(x, y, translated)
                done += 1
                if progress_cb:
                    progress_cb(int(done / total * 100))
            c.showPage()
        c.save()
    stats = (translator.total_count, translator.fail_count, translator.paid_fallback_count, translator.fallback_engine_count, translator.google_skipped)
    return "\n".join(preview_lines), stats


def translate_pdf_multi(in_path, out_path, translators, progress_cb=None):
    """Rebuild the PDF as a flowing document: each original line is followed
    by one labelled line per target language, auto-paginating (original
    absolute layout can't be preserved once the content grows 4x taller)."""
    import xml.sax.saxutils as saxutils
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
    from reportlab.lib.styles import ParagraphStyle

    style_normal = ParagraphStyle(
        "normal", fontName=UNICODE_FONT_NAME, fontSize=11, leading=14, spaceAfter=4
    )
    style_translation = ParagraphStyle(
        "translation", fontName=UNICODE_FONT_NAME, fontSize=11, leading=14, spaceAfter=2
    )

    preview_lines = []
    flowables = []
    with pdfplumber.open(in_path) as pdf:
        # Progress is reported per translated line (across all target
        # languages), not per page -- see translate_pdf for why.
        pages_lines = []
        total_lines = 0
        for page in pdf.pages:
            lines = [l for l in (page.extract_text_lines() or []) if l.get("text", "").strip()]
            pages_lines.append(lines)
            total_lines += len(lines)
        total = (total_lines * len(translators)) or 1

        done = 0
        with ThreadPoolExecutor(max_workers=max(1, len(translators))) as pool:
            for i, lines in enumerate(pages_lines):
                if i > 0:
                    flowables.append(PageBreak())
                preview_lines.append(f"--- Page {i + 1} ---")
                for line in lines:
                    text = line["text"]
                    flowables.append(Paragraph(saxutils.escape(text), style_normal))
                    preview_lines.append(text)
                    translated_map = translate_segment_all_languages(text, translators, pool)
                    for code in translators:
                        translated = translated_map[code] or text
                        label = CODE_TO_NAME[code]
                        flowables.append(Paragraph(
                            f"<b>[{label}]</b> {saxutils.escape(translated)}", style_translation
                        ))
                        preview_lines.append(f"[{label}] {translated}")
                        done += 1
                        if progress_cb:
                            progress_cb(int(done / total * 100))
                    flowables.append(Spacer(1, 8))

    doc = SimpleDocTemplate(out_path, pagesize=letter)
    doc.build(flowables)
    stats = (
        sum(t.total_count for t in translators.values()),
        sum(t.fail_count for t in translators.values()),
        sum(t.paid_fallback_count for t in translators.values()),
        sum(t.fallback_engine_count for t in translators.values()),
        any(t.google_skipped for t in translators.values()),
    )
    return "\n".join(preview_lines), stats


def _format_stats_message(stats):
    """Turn a (total, failed, paid_fallback, fallback_engine, google_skipped)
    tuple into a human status suffix, or '' if everything translated cleanly
    via Google."""
    total, failed, paid, fallback_engine, google_skipped = stats
    parts = []
    if fallback_engine:
        plural = "s" if fallback_engine != 1 else ""
        parts.append(
            f"{fallback_engine} segment{plural} used the MyMemory fallback "
            "engine (Google was rate-limiting)."
        )
    if google_skipped:
        parts.append(
            "Google kept failing early on, so it was skipped for the rest "
            "of this run to save time."
        )
    if paid:
        plural = "s" if paid != 1 else ""
        parts.append(f"{paid} segment{plural} used the paid Google Cloud API fallback.")
    if failed:
        parts.append(
            f"Warning: {failed}/{total} segment(s) could not be translated "
            "even with every fallback and were left in the original language."
        )
    return " ".join(parts)


class TranslateWorker(QThread):
    progress = pyqtSignal(int)
    finished_ok = pyqtSignal(str, str, str)
    failed = pyqtSignal(str)

    def __init__(self, in_path, source, target):
        super().__init__()
        self.in_path = in_path
        self.source = source
        self.target = target

    def run(self):
        try:
            ext = os.path.splitext(self.in_path)[1].lower()
            source_path = self.in_path

            if ext == ".doc":
                source_path = convert_doc_to_docx(self.in_path)
                ext = ".docx"

            base, _ = os.path.splitext(self.in_path)
            out_path = f"{base}_translated_{self.target}{ext}"
            translator = Translator(self.source, self.target)

            if ext == ".xlsx":
                preview, stats = translate_xlsx(source_path, out_path, translator, self.progress.emit)
            elif ext == ".docx":
                preview, stats = translate_docx(source_path, out_path, translator, self.progress.emit)
            elif ext == ".pdf":
                preview, stats = translate_pdf(source_path, out_path, translator, self.progress.emit)
            else:
                self.failed.emit(f"Unsupported file type: {ext}")
                return

            self.finished_ok.emit(out_path, preview or "", _format_stats_message(stats))
        except Exception as exc:
            self.failed.emit(f"{exc}\n{traceback.format_exc()}")


class MultiTranslateWorker(QThread):
    progress = pyqtSignal(int)
    finished_ok = pyqtSignal(str, str, str)
    failed = pyqtSignal(str)

    def __init__(self, in_path, source, target_codes):
        super().__init__()
        self.in_path = in_path
        self.source = source
        self.target_codes = target_codes

    def run(self):
        try:
            ext = os.path.splitext(self.in_path)[1].lower()
            source_path = self.in_path

            if ext == ".doc":
                source_path = convert_doc_to_docx(self.in_path)
                ext = ".docx"

            base, _ = os.path.splitext(self.in_path)
            suffix = "_".join(self.target_codes)
            out_path = f"{base}_multilang_{suffix}{ext}"
            translators = {code: Translator(self.source, code) for code in self.target_codes}

            if ext == ".xlsx":
                preview, stats = translate_xlsx_multi(source_path, out_path, translators, self.progress.emit)
            elif ext == ".docx":
                preview, stats = translate_docx_multi(source_path, out_path, translators, self.progress.emit)
            elif ext == ".pdf":
                preview, stats = translate_pdf_multi(source_path, out_path, translators, self.progress.emit)
            else:
                self.failed.emit(f"Unsupported file type: {ext}")
                return

            self.finished_ok.emit(out_path, preview or "", _format_stats_message(stats))
        except Exception as exc:
            self.failed.emit(f"{exc}\n{traceback.format_exc()}")


class TextTranslateWorker(QThread):
    finished_ok = pyqtSignal(str, str)
    failed = pyqtSignal(str)

    def __init__(self, text, source, target):
        super().__init__()
        self.text = text
        self.source = source
        self.target = target

    def run(self):
        try:
            translator = Translator(self.source, self.target)
            result = translator.translate(self.text)
            warning = ""
            if translator.fail_count:
                warning = "Translation failed after retries — showing original text."
            self.finished_ok.emit(result or "", warning)
        except Exception as exc:
            self.failed.emit(f"{exc}\n{traceback.format_exc()}")


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Document Translator")
        self.resize(800, 640)

        self.in_path = None
        self.out_path = None
        self.worker = None
        self.text_worker = None

        tabs = QTabWidget()
        self.setCentralWidget(tabs)

        doc_tab = QWidget()
        tabs.addTab(doc_tab, "Document Translator")
        layout = QVBoxLayout(doc_tab)

        # File selection
        file_box = QGroupBox("Source Document")
        file_layout = QHBoxLayout(file_box)
        self.file_label = QLabel("No file selected")
        self.file_label.setWordWrap(True)
        browse_btn = QPushButton("Browse...")
        browse_btn.clicked.connect(self.browse_file)
        file_layout.addWidget(self.file_label, stretch=1)
        file_layout.addWidget(browse_btn)
        layout.addWidget(file_box)

        # Language selection
        lang_box = QGroupBox("Languages")
        lang_layout = QHBoxLayout(lang_box)
        self.source_combo = QComboBox()
        self.target_combo = QComboBox()
        self.source_combo.addItems(LANGUAGES.keys())
        self.target_combo.addItems(LANGUAGES.keys())
        self.source_combo.setCurrentText("French")
        self.target_combo.setCurrentText("English")

        swap_btn = QToolButton()
        swap_btn.setText("<->")
        swap_btn.clicked.connect(self.swap_languages)

        lang_layout.addWidget(QLabel("From:"))
        lang_layout.addWidget(self.source_combo)
        lang_layout.addWidget(swap_btn)
        lang_layout.addWidget(QLabel("To:"))
        lang_layout.addWidget(self.target_combo)
        layout.addWidget(lang_box)

        # Actions
        action_layout = QHBoxLayout()
        self.translate_btn = QPushButton("Translate")
        self.translate_btn.clicked.connect(self.start_translation)
        self.translate_btn.setEnabled(False)
        self.save_btn = QPushButton("Save As...")
        self.save_btn.clicked.connect(self.save_output_as)
        self.save_btn.setEnabled(False)
        self.print_btn = QPushButton("Print Translated Document")
        self.print_btn.clicked.connect(self.print_output)
        self.print_btn.setEnabled(False)
        action_layout.addWidget(self.translate_btn)
        action_layout.addWidget(self.save_btn)
        action_layout.addWidget(self.print_btn)
        layout.addLayout(action_layout)

        self.progress = QProgressBar()
        layout.addWidget(self.progress)

        self.status_label = QLabel("")
        self.status_label.setWordWrap(True)
        layout.addWidget(self.status_label)

        preview_box = QGroupBox("Translated Preview")
        preview_layout = QVBoxLayout(preview_box)
        self.preview = QTextEdit()
        self.preview.setReadOnly(True)
        preview_layout.addWidget(self.preview)
        layout.addWidget(preview_box, stretch=1)

        # --- Quick Text tab ---
        text_tab = QWidget()
        tabs.addTab(text_tab, "Quick Text")
        text_layout = QVBoxLayout(text_tab)

        text_lang_box = QGroupBox("Languages")
        text_lang_layout = QHBoxLayout(text_lang_box)
        self.text_source_combo = QComboBox()
        self.text_target_combo = QComboBox()
        self.text_source_combo.addItems(LANGUAGES.keys())
        self.text_target_combo.addItems(LANGUAGES.keys())
        self.text_source_combo.setCurrentText("French")
        self.text_target_combo.setCurrentText("English")

        text_swap_btn = QToolButton()
        text_swap_btn.setText("<->")
        text_swap_btn.clicked.connect(self.swap_text_languages)

        text_lang_layout.addWidget(QLabel("From:"))
        text_lang_layout.addWidget(self.text_source_combo)
        text_lang_layout.addWidget(text_swap_btn)
        text_lang_layout.addWidget(QLabel("To:"))
        text_lang_layout.addWidget(self.text_target_combo)
        text_layout.addWidget(text_lang_box)

        input_box = QGroupBox("Paste text here")
        input_layout = QVBoxLayout(input_box)
        self.text_input = QTextEdit()
        input_layout.addWidget(self.text_input)
        text_layout.addWidget(input_box, stretch=1)

        text_action_layout = QHBoxLayout()
        self.text_translate_btn = QPushButton("Translate")
        self.text_translate_btn.clicked.connect(self.start_text_translation)
        self.text_copy_btn = QPushButton("Copy Result")
        self.text_copy_btn.clicked.connect(self.copy_text_result)
        self.text_save_btn = QPushButton("Save As...")
        self.text_save_btn.clicked.connect(self.save_text_result_as)
        self.text_print_btn = QPushButton("Print Result")
        self.text_print_btn.clicked.connect(self.print_text_result)
        text_action_layout.addWidget(self.text_translate_btn)
        text_action_layout.addWidget(self.text_copy_btn)
        text_action_layout.addWidget(self.text_save_btn)
        text_action_layout.addWidget(self.text_print_btn)
        text_layout.addLayout(text_action_layout)

        self.text_status_label = QLabel("")
        self.text_status_label.setWordWrap(True)
        text_layout.addWidget(self.text_status_label)

        output_box = QGroupBox("Translated result")
        output_layout = QVBoxLayout(output_box)
        self.text_output = QTextEdit()
        self.text_output.setReadOnly(True)
        output_layout.addWidget(self.text_output)
        text_layout.addWidget(output_box, stretch=1)

        # --- Multi-Language tab ---
        multi_tab = QWidget()
        tabs.addTab(multi_tab, "Multi-Language")
        multi_layout = QVBoxLayout(multi_tab)

        self.multi_in_path = None
        self.multi_out_path = None
        self.multi_worker = None

        multi_file_box = QGroupBox("Source Document")
        multi_file_layout = QHBoxLayout(multi_file_box)
        self.multi_file_label = QLabel("No file selected")
        self.multi_file_label.setWordWrap(True)
        multi_browse_btn = QPushButton("Browse...")
        multi_browse_btn.clicked.connect(self.browse_multi_file)
        multi_file_layout.addWidget(self.multi_file_label, stretch=1)
        multi_file_layout.addWidget(multi_browse_btn)
        multi_layout.addWidget(multi_file_box)

        multi_lang_box = QGroupBox("Languages")
        multi_lang_layout = QVBoxLayout(multi_lang_box)
        source_row = QHBoxLayout()
        source_row.addWidget(QLabel("Source language:"))
        self.multi_source_combo = QComboBox()
        self.multi_source_combo.addItems(LANGUAGES.keys())
        self.multi_source_combo.setCurrentText("English")
        source_row.addWidget(self.multi_source_combo)
        source_row.addStretch(1)
        multi_lang_layout.addLayout(source_row)

        multi_lang_layout.addWidget(QLabel("Translate into (each paragraph will show all selected languages, stacked):"))
        checkbox_row = QHBoxLayout()
        self.multi_target_checkboxes = {}
        for name, code in LANGUAGES.items():
            cb = QCheckBox(name)
            self.multi_target_checkboxes[code] = cb
            checkbox_row.addWidget(cb)
        self.multi_target_checkboxes["hi"].setChecked(True)
        self.multi_target_checkboxes["bn"].setChecked(True)
        self.multi_target_checkboxes["mg"].setChecked(True)
        multi_lang_layout.addLayout(checkbox_row)
        multi_layout.addWidget(multi_lang_box)

        multi_action_layout = QHBoxLayout()
        self.multi_translate_btn = QPushButton("Translate")
        self.multi_translate_btn.clicked.connect(self.start_multi_translation)
        self.multi_translate_btn.setEnabled(False)
        self.multi_save_btn = QPushButton("Save As...")
        self.multi_save_btn.clicked.connect(self.save_multi_output_as)
        self.multi_save_btn.setEnabled(False)
        self.multi_print_btn = QPushButton("Print Translated Document")
        self.multi_print_btn.clicked.connect(self.print_multi_output)
        self.multi_print_btn.setEnabled(False)
        multi_action_layout.addWidget(self.multi_translate_btn)
        multi_action_layout.addWidget(self.multi_save_btn)
        multi_action_layout.addWidget(self.multi_print_btn)
        multi_layout.addLayout(multi_action_layout)

        self.multi_progress = QProgressBar()
        multi_layout.addWidget(self.multi_progress)

        self.multi_status_label = QLabel("")
        self.multi_status_label.setWordWrap(True)
        multi_layout.addWidget(self.multi_status_label)

        multi_preview_box = QGroupBox("Translated Preview")
        multi_preview_layout = QVBoxLayout(multi_preview_box)
        self.multi_preview = QTextEdit()
        self.multi_preview.setReadOnly(True)
        multi_preview_layout.addWidget(self.multi_preview)
        multi_layout.addWidget(multi_preview_box, stretch=1)

    def swap_languages(self):
        s, t = self.source_combo.currentText(), self.target_combo.currentText()
        self.source_combo.setCurrentText(t)
        self.target_combo.setCurrentText(s)

    def browse_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select document", "",
            "Documents (*.xlsx *.docx *.doc *.pdf)"
        )
        if path:
            self.in_path = path
            self.file_label.setText(path)
            self.translate_btn.setEnabled(True)
            self.save_btn.setEnabled(False)
            self.print_btn.setEnabled(False)
            self.out_path = None

    def start_translation(self):
        if not self.in_path:
            return
        source = LANGUAGES[self.source_combo.currentText()]
        target = LANGUAGES[self.target_combo.currentText()]
        if source == target:
            QMessageBox.warning(self, "Same language", "Source and target languages must differ.")
            return

        self.translate_btn.setEnabled(False)
        self.print_btn.setEnabled(False)
        self.progress.setValue(0)
        self.preview.clear()
        self.status_label.setText(
            f"Translating {os.path.basename(self.in_path)} ({source} -> {target})..."
        )

        self.worker = TranslateWorker(self.in_path, source, target)
        self.worker.progress.connect(self.progress.setValue)
        self.worker.finished_ok.connect(self.on_translation_done)
        self.worker.failed.connect(self.on_translation_failed)
        self.worker.start()

    def on_translation_done(self, out_path, preview_text, stats_msg):
        self.out_path = out_path
        self.translate_btn.setEnabled(True)
        self.save_btn.setEnabled(True)
        self.print_btn.setEnabled(True)
        self.preview.setPlainText(preview_text)
        status = f"Done. Saved to: {out_path}"
        if stats_msg:
            status += "  " + stats_msg
        self.status_label.setText(status)

    def on_translation_failed(self, message):
        self.translate_btn.setEnabled(True)
        self.status_label.setText(f"ERROR: {message.splitlines()[0]}")
        QMessageBox.critical(self, "Translation failed", message.split("\n")[0])

    def save_output_as(self):
        if not self.out_path or not os.path.exists(self.out_path):
            return
        ext = os.path.splitext(self.out_path)[1]
        filters = {
            ".xlsx": "Excel Files (*.xlsx)",
            ".docx": "Word Files (*.docx)",
            ".pdf": "PDF Files (*.pdf)",
        }.get(ext, "All Files (*)")
        default_name = os.path.basename(self.out_path)
        dest, _ = QFileDialog.getSaveFileName(self, "Save translated document as", default_name, filters)
        if dest:
            try:
                import shutil
                shutil.copyfile(self.out_path, dest)
                self.status_label.setText(f"Saved to: {dest}")
            except Exception as exc:
                QMessageBox.critical(self, "Save failed", str(exc))

    def print_output(self):
        if not self.out_path or not os.path.exists(self.out_path):
            return
        try:
            if sys.platform == "win32":
                os.startfile(self.out_path, "print")
                self.status_label.setText(f"Sent to printer: {self.out_path}")
            else:
                QMessageBox.information(
                    self, "Print",
                    f"Open and print this file manually: {self.out_path}"
                )
        except Exception as exc:
            QMessageBox.critical(self, "Print failed", str(exc))

    # --- Quick Text tab ---

    def swap_text_languages(self):
        s, t = self.text_source_combo.currentText(), self.text_target_combo.currentText()
        self.text_source_combo.setCurrentText(t)
        self.text_target_combo.setCurrentText(s)

    def start_text_translation(self):
        text = self.text_input.toPlainText()
        if not text.strip():
            return
        source = LANGUAGES[self.text_source_combo.currentText()]
        target = LANGUAGES[self.text_target_combo.currentText()]
        if source == target:
            QMessageBox.warning(self, "Same language", "Source and target languages must differ.")
            return

        self.text_translate_btn.setEnabled(False)
        self.text_status_label.setText(f"Translating ({source} -> {target})...")

        self.text_worker = TextTranslateWorker(text, source, target)
        self.text_worker.finished_ok.connect(self.on_text_translation_done)
        self.text_worker.failed.connect(self.on_text_translation_failed)
        self.text_worker.start()

    def on_text_translation_done(self, result, warning):
        self.text_translate_btn.setEnabled(True)
        self.text_output.setPlainText(result)
        self.text_status_label.setText(warning if warning else "Done.")

    def on_text_translation_failed(self, message):
        self.text_translate_btn.setEnabled(True)
        self.text_status_label.setText(f"ERROR: {message.splitlines()[0]}")
        QMessageBox.critical(self, "Translation failed", message.split("\n")[0])

    def copy_text_result(self):
        QApplication.clipboard().setText(self.text_output.toPlainText())
        self.text_status_label.setText("Copied to clipboard.")

    def save_text_result_as(self):
        text = self.text_output.toPlainText()
        if not text.strip():
            return
        dest, _ = QFileDialog.getSaveFileName(
            self, "Save translated text as", "translated.txt", "Text Files (*.txt)"
        )
        if dest:
            try:
                with open(dest, "w", encoding="utf-8") as f:
                    f.write(text)
                self.text_status_label.setText(f"Saved to: {dest}")
            except Exception as exc:
                QMessageBox.critical(self, "Save failed", str(exc))

    def print_text_result(self):
        text = self.text_output.toPlainText()
        if not text.strip():
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        dialog = QPrintDialog(printer, self)
        if dialog.exec() == QPrintDialog.DialogCode.Accepted:
            self.text_output.print_(printer)
            self.text_status_label.setText("Sent to printer.")

    # --- Multi-Language tab ---

    def browse_multi_file(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "Select document", "",
            "Documents (*.xlsx *.docx *.doc *.pdf)"
        )
        if path:
            self.multi_in_path = path
            self.multi_file_label.setText(path)
            self.multi_translate_btn.setEnabled(True)
            self.multi_save_btn.setEnabled(False)
            self.multi_print_btn.setEnabled(False)
            self.multi_out_path = None

    def start_multi_translation(self):
        if not self.multi_in_path:
            return
        source = LANGUAGES[self.multi_source_combo.currentText()]
        target_codes = [code for code, cb in self.multi_target_checkboxes.items() if cb.isChecked()]
        target_codes = [code for code in target_codes if code != source]
        if not target_codes:
            QMessageBox.warning(self, "No target languages", "Select at least one target language (different from the source).")
            return

        self.multi_translate_btn.setEnabled(False)
        self.multi_save_btn.setEnabled(False)
        self.multi_print_btn.setEnabled(False)
        self.multi_progress.setValue(0)
        self.multi_preview.clear()
        lang_names = ", ".join(CODE_TO_NAME[c] for c in target_codes)
        self.multi_status_label.setText(
            f"Translating {os.path.basename(self.multi_in_path)} into {lang_names}..."
        )

        self.multi_worker = MultiTranslateWorker(self.multi_in_path, source, target_codes)
        self.multi_worker.progress.connect(self.multi_progress.setValue)
        self.multi_worker.finished_ok.connect(self.on_multi_translation_done)
        self.multi_worker.failed.connect(self.on_multi_translation_failed)
        self.multi_worker.start()

    def on_multi_translation_done(self, out_path, preview_text, stats_msg):
        self.multi_out_path = out_path
        self.multi_translate_btn.setEnabled(True)
        self.multi_save_btn.setEnabled(True)
        self.multi_print_btn.setEnabled(True)
        self.multi_preview.setPlainText(preview_text)
        status = f"Done. Saved to: {out_path}"
        if stats_msg:
            status += "  " + stats_msg
        self.multi_status_label.setText(status)

    def on_multi_translation_failed(self, message):
        self.multi_translate_btn.setEnabled(True)
        self.multi_status_label.setText(f"ERROR: {message.splitlines()[0]}")
        QMessageBox.critical(self, "Translation failed", message.split("\n")[0])

    def save_multi_output_as(self):
        if not self.multi_out_path or not os.path.exists(self.multi_out_path):
            return
        ext = os.path.splitext(self.multi_out_path)[1]
        filters = {
            ".xlsx": "Excel Files (*.xlsx)",
            ".docx": "Word Files (*.docx)",
            ".pdf": "PDF Files (*.pdf)",
        }.get(ext, "All Files (*)")
        default_name = os.path.basename(self.multi_out_path)
        dest, _ = QFileDialog.getSaveFileName(self, "Save translated document as", default_name, filters)
        if dest:
            try:
                import shutil
                shutil.copyfile(self.multi_out_path, dest)
                self.multi_status_label.setText(f"Saved to: {dest}")
            except Exception as exc:
                QMessageBox.critical(self, "Save failed", str(exc))

    def print_multi_output(self):
        if not self.multi_out_path or not os.path.exists(self.multi_out_path):
            return
        try:
            if sys.platform == "win32":
                os.startfile(self.multi_out_path, "print")
                self.multi_status_label.setText(f"Sent to printer: {self.multi_out_path}")
            else:
                QMessageBox.information(
                    self, "Print",
                    f"Open and print this file manually: {self.multi_out_path}"
                )
        except Exception as exc:
            QMessageBox.critical(self, "Print failed", str(exc))


def main():
    app = QApplication(sys.argv)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
